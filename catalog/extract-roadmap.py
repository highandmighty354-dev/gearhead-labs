#!/usr/bin/env python3
"""Gearhead Labs — Premium roadmap extractor.

Reads the owner's authoritative workbook catalog/source/Gearhead_Labs_Premium_Master_Roadmap.xlsx
(tabs: "Premium Roadmap 140", "New Research Additions", "Do Not Add") and writes
catalog/premium-roadmap.json: every roadmap row with a stable id, its workbook reference and decision.
Nothing is typed by hand; ids are derived from the tool name and never reused.

Usage:  python3 catalog/extract-roadmap.py            write catalog/premium-roadmap.json
        python3 catalog/extract-roadmap.py --check    fail if the committed file differs
Needs openpyxl.
"""
import hashlib, json, os, re, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'source', 'Gearhead_Labs_Premium_Master_Roadmap.xlsx')
OUT = os.path.join(HERE, 'premium-roadmap.json')

def slug(name):
    s = name.lower().replace('→', ' to ').replace('–', ' ').replace('/', ' ').replace('+', ' and ').replace("'", '')
    return 'p_' + re.sub(r'[^a-z0-9]+', '_', s).strip('_')

def kind(t):
    t = (t or '').strip().lower()
    return {'calculator': 'calculator', 'analyzer': 'analyzer', 'workbench': 'workbench'}[t]

def build():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb['Premium Roadmap 140']
    rows = list(ws.iter_rows(values_only=True))
    summary = {str(r[0]).strip(): r[1] for r in rows[1:9] if r and r[0]}
    head = [str(h).strip() if h else '' for h in rows[10]]
    assert head[:8] == ['Source', 'Rank', 'Section', 'Calculator / Tool', 'Type', 'Original Domain', 'Decision', 'Reason'], head
    items, excluded = [], []
    for n, r in enumerate(rows[11:], start=12):
        if not r or not r[3]: continue
        src, rank, section, name, typ, domain, decision, reason = (r + (None,) * 8)[:8]
        rec = {'id': slug(name), 'name': str(name).strip(), 'section': str(section).strip(), 'kind': kind(typ),
               'origin': {'sheet': 'Premium Roadmap 140', 'row': n, 'list': str(src).strip(), 'rank': rank,
                          'domain': (str(domain).strip() if domain else None)},
               'workbook_decision': str(decision).strip()}
        if str(decision).startswith('DO NOT ADD'):
            rec['reason'] = str(reason).strip(); excluded.append(rec)
        else:
            items.append(rec)
    ws = wb['New Research Additions']
    for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or not r[1]: continue
        section, name, typ, why = (r + (None,) * 4)[:4]
        items.append({'id': slug(name), 'name': str(name).strip(), 'section': str(section).strip(), 'kind': kind(typ),
                      'origin': {'sheet': 'New Research Additions', 'row': n, 'list': 'New Research Additions', 'rank': None, 'domain': None},
                      'workbook_decision': 'NEW RESEARCH ADDITION', 'note': str(why).strip()})
    dna = [str(r[0]).strip() for r in wb['Do Not Add'].iter_rows(min_row=2, values_only=True) if r and r[0]]
    ids = [i['id'] for i in items + excluded]
    assert len(ids) == len(set(ids)), 'duplicate roadmap ids'
    assert sorted(dna) == sorted(e['name'] for e in excluded), 'Do Not Add tab differs from the excluded roadmap rows'
    with open(SRC, 'rb') as f: sha = hashlib.sha256(f.read()).hexdigest()
    return {
        'schema': 'gearhead-premium-roadmap/1',
        'source': {'file': 'catalog/source/Gearhead_Labs_Premium_Master_Roadmap.xlsx', 'sha256': sha},
        'workbook_summary': {k: v for k, v in summary.items()},
        'counts': {'historical_recovered': len(items) + len(excluded) - sum(1 for i in items if i['origin']['sheet'] == 'New Research Additions'),
                   'do_not_add': len(excluded),
                   'historical_retained': sum(1 for i in items if i['origin']['sheet'] == 'Premium Roadmap 140'),
                   'new_research_additions': sum(1 for i in items if i['origin']['sheet'] == 'New Research Additions'),
                   'working_pool': len(items)},
        'candidates': items,
        'do_not_add': excluded,
    }

def serialize(d):
    head = {k: v for k, v in d.items() if k not in ('candidates', 'do_not_add')}
    txt = json.dumps(head, indent=2, ensure_ascii=False)[:-2]
    line = lambda x: '    ' + json.dumps(x, ensure_ascii=False)
    return (txt + ',\n  "candidates": [\n' + ',\n'.join(map(line, d['candidates'])) + '\n  ],\n  "do_not_add": [\n'
            + ',\n'.join(map(line, d['do_not_add'])) + '\n  ]\n}\n')

if __name__ == '__main__':
    text = serialize(build())
    if '--check' in sys.argv:
        same = os.path.exists(OUT) and open(OUT, encoding='utf-8').read() == text
        print('premium roadmap: committed file matches the workbook' if same else 'premium roadmap: OUT OF DATE (run python3 catalog/extract-roadmap.py)')
        sys.exit(0 if same else 1)
    open(OUT, 'w', encoding='utf-8').write(text)
    print('written catalog/premium-roadmap.json')
